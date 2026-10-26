#!/usr/bin/env python
# coding: utf-8

# In[1]:


from pathlib import Path
from dotenv import load_dotenv
import os

# Find the project root directory
BASE_DIR = Path(__file__).resolve().parent.parent

# Locate .env in the project root
env_path = BASE_DIR / ".env"

# Load environment variables
load_dotenv(dotenv_path=env_path)

# Verify configuration
print("Project directory:", BASE_DIR)
print("Environment file exists:", env_path.exists())
print("Database host:", os.getenv("DB_HOST"))
print("Database name:", os.getenv("DB_NAME"))
print("Database user:", os.getenv("DB_USER"))
print("Password configured:", bool(os.getenv("DB_PASSWORD")))
print(".gitignore exists:", (BASE_DIR / ".gitignore").exists())


# In[11]:



# ============================================================
# PROJECT: AUTOMATED LOGISTICS KPI REPORTING SYSTEM
# FILE: generate_reports.py
#
# PURPOSE:
# 1. Connect to MySQL
# 2. Execute logistics KPI queries
# 3. Convert SQL results into Pandas DataFrames
# 4. Export a multi-sheet Excel report
# 5. Log execution status
# ============================================================

import logging
import os
from datetime import datetime
from pathlib import Path

import mysql.connector
import pandas as pd
from dotenv import load_dotenv
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# 1. PROJECT CONFIGURATION
# ============================================================

# Locate the project root
BASE_DIR = Path(__file__).resolve().parent.parent

# Define output directories
REPORT_DIR = BASE_DIR / "reports"
LOG_DIR = BASE_DIR / "logs"

# Create directories if they don't exist
REPORT_DIR.mkdir(parents=True, exist_ok=True)
LOG_DIR.mkdir(parents=True, exist_ok=True)

load_dotenv(BASE_DIR / ".env")

logging.basicConfig(
    filename=LOG_DIR / "reporting.log",
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)


# ============================================================
# 2. MYSQL DATABASE CONNECTION
# ============================================================

def connect_database():
    """Establish a connection to the MySQL database."""

    connection = mysql.connector.connect(
        host=os.getenv("DB_HOST", "localhost"),
        port=int(os.getenv("DB_PORT", "3306")),
        user=os.environ["DB_USER"],
        password=os.environ["DB_PASSWORD"],
        database=os.environ["DB_NAME"]
    )

    print("Successfully connected to MySQL!")
    return connection


# ============================================================
# 3. SQL KPI QUERIES
# ============================================================

KPI_QUERIES = {

    # KPI 1: Overall on-time delivery
    "Overall_KPIs": """
        SELECT
            COUNT(*) AS completed_orders,
            SUM(
                CASE
                    WHEN delivery_status IN (
                        'Shipping on time',
                        'Advance shipping'
                    )
                    THEN 1 ELSE 0
                END
            ) AS on_time_orders,
            ROUND(
                100.0 * AVG(
                    CASE
                        WHEN delivery_status IN (
                            'Shipping on time',
                            'Advance shipping'
                        )
                        THEN 1.0 ELSE 0.0
                    END
                ), 2
            ) AS on_time_delivery_pct
        FROM orders
        WHERE delivery_status <> 'Shipping canceled'
    """,

    # KPI 2: Shipping mode performance
    "Shipping_Modes": """
        SELECT
            shipping_mode,
            COUNT(*) AS total_orders,
            ROUND(AVG(actual_shipping_days), 2)
                AS avg_actual_shipping_days,
            ROUND(AVG(scheduled_shipping_days), 2)
                AS avg_scheduled_shipping_days,
            ROUND(
                100.0 * AVG(
                    CASE
                        WHEN delivery_status IN (
                            'Shipping on time',
                            'Advance shipping'
                        )
                        THEN 1.0 ELSE 0.0
                    END
                ), 2
            ) AS on_time_delivery_pct
        FROM orders
        WHERE delivery_status <> 'Shipping canceled'
        GROUP BY shipping_mode
        ORDER BY on_time_delivery_pct DESC
    """,

    # KPI 3: Regional delivery performance
    "Regional_KPIs": """
        SELECT
            order_region,
            COUNT(*) AS total_orders,
            SUM(
                CASE
                    WHEN delivery_status = 'Late delivery'
                    THEN 1 ELSE 0
                END
            ) AS late_orders,
            ROUND(
                100.0 * AVG(
                    CASE
                        WHEN delivery_status = 'Late delivery'
                        THEN 1.0 ELSE 0.0
                    END
                ), 2
            ) AS late_delivery_pct
        FROM orders
        WHERE delivery_status <> 'Shipping canceled'
        GROUP BY order_region
        HAVING COUNT(*) >= 100
        ORDER BY late_delivery_pct DESC
    """,

    # KPI 4: Sales exposure by delivery status
    "Sales_Exposure": """
        SELECT
            o.delivery_status,
            COUNT(DISTINCT o.order_id) AS total_orders,
            ROUND(SUM(oi.sales), 2) AS total_sales
        FROM orders o
        INNER JOIN order_items oi
            ON o.order_id = oi.order_id
        GROUP BY o.delivery_status
        ORDER BY total_sales DESC
    """,

    # KPI 5: Shipping schedule variance
    "Schedule_Variance": """
        SELECT
            shipping_mode,
            COUNT(*) AS total_orders,
            ROUND(
                AVG(
                    actual_shipping_days -
                    scheduled_shipping_days
                ), 2
            ) AS avg_schedule_variance,
            SUM(
                CASE
                    WHEN actual_shipping_days >
                         scheduled_shipping_days
                    THEN 1 ELSE 0
                END
            ) AS orders_exceeding_schedule
        FROM orders
        WHERE delivery_status <> 'Shipping canceled'
        GROUP BY shipping_mode
        ORDER BY avg_schedule_variance DESC
    """,

    # KPI 6: Monthly delivery performance
    "Monthly_Trends": """
        SELECT
            DATE_FORMAT(order_date, '%Y-%m')
                AS order_month,
            COUNT(*) AS completed_orders,
            SUM(
                CASE
                    WHEN delivery_status = 'Late delivery'
                    THEN 1 ELSE 0
                END
            ) AS late_orders,
            ROUND(
                100.0 * AVG(
                    CASE
                        WHEN delivery_status IN (
                            'Shipping on time',
                            'Advance shipping'
                        )
                        THEN 1.0 ELSE 0.0
                    END
                ), 2
            ) AS on_time_delivery_pct
        FROM orders
        WHERE delivery_status <> 'Shipping canceled'
        GROUP BY DATE_FORMAT(order_date, '%Y-%m')
        ORDER BY order_month
    """
}


# ============================================================
# 4. EXECUTE SQL AND RETURN DATAFRAME
# ============================================================

def run_query(connection, query):
    """Execute SQL and return results as a DataFrame."""

    cursor = connection.cursor()

    try:
        cursor.execute(query)

        columns = [
            column[0]
            for column in cursor.description
        ]

        rows = cursor.fetchall()

        # Convert SQL results into a Pandas DataFrame
        df = pd.DataFrame(rows, columns=columns)

        # Convert numeric SQL results into proper numeric types
        for column in df.columns:
            if column.endswith((
                "_pct",
                "_sales",
                "_days",
                "_variance",
                "_orders",
                "_schedule"
            )):
                df[column] = pd.to_numeric(
                    df[column],
                    errors="raise"
                )

        return df

    finally:
        cursor.close()



# ============================================================
# 5. FORMAT EXCEL WORKSHEETS
# ============================================================

def format_worksheet(worksheet):
    """Apply basic professional formatting."""

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="17365D"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center"
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column in worksheet.columns:
        max_length = max(
            len(str(cell.value))
            if cell.value is not None else 0
            for cell in column
        )

        column_letter = get_column_letter(
            column[0].column
        )

        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 3, 35)


# ============================================================
# 6. GENERATE EXCEL KPI REPORT
# ============================================================

def generate_report(connection):
    """Run all KPI queries and export an Excel workbook."""

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    output_file = (
        REPORT_DIR /
        f"logistics_kpi_report_{timestamp}.xlsx"
    )

    with pd.ExcelWriter(
        output_file,
        engine="openpyxl"
    ) as writer:

        for sheet_name, query in KPI_QUERIES.items():

            print(f"Generating: {sheet_name}")

            df = run_query(connection, query)

            df.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False
            )

            worksheet = writer.sheets[sheet_name]
            format_worksheet(worksheet)

            logging.info(
                "%s generated with %s rows",
                sheet_name,
                len(df)
            )

    print(f"\nReport saved to: {output_file}")

    logging.info(
        "Report successfully generated: %s",
        output_file
    )

    return output_file


# ============================================================
# 7. MAIN AUTOMATION WORKFLOW
# ============================================================

def main():

    connection = None

    try:
        logging.info("Starting KPI reporting process")

        connection = connect_database()

        generate_report(connection)

        print("\nAutomation completed successfully!")

    except Exception as error:

        logging.exception("Reporting automation failed")
        print(f"Error: {error}")
        raise

    finally:

        if connection is not None and connection.is_connected():
            connection.close()
            print("MySQL connection closed.")


# ============================================================
# 8. SCRIPT ENTRY POINT
# ============================================================

if __name__ == "__main__":
    main()


# In[ ]:





# In[12]:


import pandas as pd
from pathlib import Path

# Project root directory
BASE_DIR = Path(__file__).resolve().parent.parent

# Reports folder inside the project
report_dir = BASE_DIR / "reports"

# Find generated Excel reports
report_files = sorted(
    report_dir.glob("logistics_kpi_report_*.xlsx")
)

if not report_files:
    raise FileNotFoundError(f"No Excel reports found in {report_dir}")

latest_report = report_files[-1]

excel_file = pd.ExcelFile(latest_report)

print("Worksheets:", excel_file.sheet_names)

overall_kpis = pd.read_excel(
    latest_report,
    sheet_name="Overall_KPIs"
)

print(overall_kpis)
# In[ ]:




